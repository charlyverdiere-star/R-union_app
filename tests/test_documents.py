"""Génération des livrables 2 et 3.

Le point le plus important : le rapport complet et la synthèse sont rendus
depuis le **même** `ReportModel`. Ils ne peuvent donc pas se contredire, et
c'est exactement ce que vérifie `test_rapport_et_synthese_sont_coherents`.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from backend.app.models import (
    ActionItem, Conclusion, Decision, Evidence, ExecutiveSummary, KeyPoint,
    MeetingMeta, OpenQuestion, ReportModel, RiskItem, Segment, TopicSection,
    Transcript, Verification,
)
from backend.app.render import common, docx_renderer, pdf_renderer
from backend.app.storage import deliverable_name, slugify


def texte_docx(chemin: Path) -> str:
    document = Document(chemin)
    morceaux = [p.text for p in document.paragraphs]
    for tableau in document.tables:
        for ligne in tableau.rows:
            morceaux.extend(cellule.text for cellule in ligne.cells)
    return "\n".join(morceaux)


def texte_pdf(chemin: Path) -> str:
    pdfium = pytest.importorskip("pypdfium2")
    document = pdfium.PdfDocument(chemin)
    return "\n".join(document[i].get_textpage().get_text_range() for i in range(len(document)))


# --- nommage ----------------------------------------------------------------


def test_nom_de_fichier_suit_la_convention():
    nom = deliverable_name("2026-10-15", "Réunion d'équipe — Été", "report", ".docx")
    assert nom == "2026-10-15_Reunion-d-equipe-Ete_rapport-complet.docx"


def test_nom_de_fichier_sans_accent_ni_caractere_special():
    nom = deliverable_name("2026-01-02", "Café & Co : bilan 100 %", "summary", ".pdf")
    assert re.fullmatch(r"[A-Za-z0-9_.\-]+", nom), f"nom non portable : {nom}"


def test_titre_vide_reste_exploitable():
    assert slugify("???") == "reunion"


# --- contenu du rapport complet ---------------------------------------------


def test_rapport_contient_les_sept_sections(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    cible = docx_renderer.render_report(modele, transcription, tmp_path / "r.docx")
    contenu = texte_docx(cible)
    for section in [
        "1. Informations générales", "2. Résumé exécutif", "3. Compte rendu détaillé",
        "4. Décisions prises", "5. Plan d'action",
        "6. Questions ouvertes et points à clarifier", "7. Conclusion",
    ]:
        assert section in contenu, f"section manquante : {section}"


def test_rapport_porte_des_horodatages(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert len(re.findall(r"\[\d{2}:\d{2}:\d{2}\]", contenu)) >= 10, (
        "chaque information doit pouvoir être retrouvée dans l'audio"
    )


def test_rapport_signale_les_informations_a_verifier(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert common.TO_VERIFY in contenu
    assert "n'a pas pu être retrouvée mot pour mot" in contenu, (
        "la signification du marqueur doit être expliquée dans le document"
    )


def test_rapport_annonce_ses_limites(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert "Limites de ce compte rendu" in contenu
    assert "écartées" in contenu, "les informations retirées doivent être annoncées"


def test_rapport_distingue_action_ferme_et_proposee(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert "engagement ferme" in contenu or "action proposée" in contenu


def test_absence_de_responsable_explicite(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert "non precise" in contenu or "non précisé" in contenu


# --- synthèse ----------------------------------------------------------------


def test_synthese_est_courte_et_actionnable(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_summary(modele, transcription, tmp_path / "s.docx"))
    for rubrique in ["L'essentiel", "Décisions importantes", "Actions à retenir",
                     "Points de vigilance", "Questions en suspens"]:
        assert rubrique in contenu, f"rubrique manquante : {rubrique}"
    rapport = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert len(contenu) < len(rapport) * 0.6, "la synthèse doit rester nettement plus courte"


def test_synthese_affiche_des_libelles_lisibles(reunion_analysee, tmp_path):
    """Les codes internes du modèle de données ne doivent jamais apparaître
    dans un document destiné à être diffusé."""
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_summary(modele, transcription, tmp_path / "s.docx"))
    for code in ("(difficulte)", "(dependance)", "(probleme)",
                 "(sans reponse)", "(a verifier)", "(a reprendre)",
                 "non precise", "non precisee"):
        assert code not in contenu, f"code interne affiché tel quel : {code}"
    assert "dépendance" in contenu or "difficulté" in contenu or "risque" in contenu


def test_documents_datent_en_francais(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    contenu = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    assert "Généré le" in contenu
    assert re.search(r"Généré le \d{2}/\d{2}/\d{4} à \d{2}:\d{2}", contenu), (
        "la date de génération doit être lisible en français"
    )


# --- cohérence entre les deux documents --------------------------------------


def test_rapport_et_synthese_sont_coherents(reunion_analysee, tmp_path):
    """Exigence centrale : aucune contradiction entre les deux livrables.

    Les deux documents étant rendus depuis le même objet, chaque décision et
    chaque action de la synthèse doit se retrouver à l'identique dans le
    rapport complet.
    """
    _, transcription, modele, _ = reunion_analysee
    rapport = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    synthese = texte_docx(docx_renderer.render_summary(modele, transcription, tmp_path / "s.docx"))

    for decision in modele.decisions:
        assert decision.statement in rapport, "décision absente du rapport"
        assert decision.statement in synthese, "décision absente de la synthèse"

    for action in modele.actions:
        assert action.task in rapport
        assert action.task in synthese
        if action.owner:
            assert action.owner in synthese, "responsable incohérent entre les documents"


def test_les_deux_formats_portent_le_meme_contenu(reunion_analysee, tmp_path):
    _, transcription, modele, _ = reunion_analysee
    docx = texte_docx(docx_renderer.render_report(modele, transcription, tmp_path / "r.docx"))
    pdf = texte_pdf(pdf_renderer.render_report(modele, transcription, tmp_path / "r.pdf"))
    for decision in modele.decisions:
        debut = decision.statement[:45]
        assert debut in docx
        assert debut.replace("'", "'") in pdf or debut in pdf, "décision absente du PDF"


# --- PDF ---------------------------------------------------------------------


def test_pdf_genere_plusieurs_pages(reunion_analysee, tmp_path):
    pdfium = pytest.importorskip("pypdfium2")
    _, transcription, modele, _ = reunion_analysee
    cible = pdf_renderer.render_report(modele, transcription, tmp_path / "r.pdf")
    assert cible.stat().st_size > 10_000
    assert len(pdfium.PdfDocument(cible)) >= 3


def test_pdf_synthese_tient_en_peu_de_pages(reunion_analysee, tmp_path):
    pdfium = pytest.importorskip("pypdfium2")
    _, transcription, modele, _ = reunion_analysee
    cible = pdf_renderer.render_summary(modele, transcription, tmp_path / "s.pdf")
    assert len(pdfium.PdfDocument(cible)) <= 4, "la synthèse doit se lire en deux minutes"


# --- cas limites --------------------------------------------------------------


@pytest.fixture
def rapport_vide():
    return ReportModel(
        meta=MeetingMeta(title="Réunion sans contenu", date="2026-10-15",
                         duration_label="2 min", duration_seconds=120)
    )


def test_rapport_vide_reste_lisible(rapport_vide, tmp_path):
    """Une réunion sans décision ne doit pas produire un document trompeur."""
    contenu = texte_docx(docx_renderer.render_report(rapport_vide, None, tmp_path / "r.docx"))
    assert contenu.count("Aucun élément de ce type") >= 4
    assert "Aucun participant n'a pu être identifié" in contenu


def test_rapport_vide_genere_aussi_en_pdf(rapport_vide, tmp_path):
    cible = pdf_renderer.render_report(rapport_vide, None, tmp_path / "r.pdf")
    assert cible.exists() and cible.stat().st_size > 1000


def test_caracteres_speciaux_ne_cassent_pas_le_pdf(tmp_path):
    """Un « & » ou un « < » venant de la transcription doit être échappé."""
    modele = ReportModel(
        meta=MeetingMeta(title="R & D <projet>", date="2026-10-15",
                         duration_label="1 h", duration_seconds=3600),
        decisions=[Decision(id="D1", statement="Migration de A & B vers <v2>",
                            evidence=Evidence(segment_ids=["s0001"]))],
    )
    transcription = Transcript(
        language="fr", provider="x", model="x", duration=3600,
        segments=[Segment(id="s0001", start=0, end=5, text="Migration de A & B vers <v2>")],
    )
    cible = pdf_renderer.render_report(modele, transcription, tmp_path / "r.pdf")
    assert cible.exists() and cible.stat().st_size > 1000
    assert "A & B" in texte_pdf(cible)


def test_horodatage_absent_si_segment_inconnu(tmp_path):
    modele = ReportModel(
        meta=MeetingMeta(title="T", date="2026-10-15", duration_label="1 h", duration_seconds=3600),
        decisions=[Decision(id="D1", statement="Une décision",
                            evidence=Evidence(segment_ids=["s9999"]))],
    )
    contenu = texte_docx(docx_renderer.render_report(modele, None, tmp_path / "r.docx"))
    assert "Une décision" in contenu
    assert "[s9999]" not in contenu
