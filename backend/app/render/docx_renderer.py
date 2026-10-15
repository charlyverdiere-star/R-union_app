"""Génération des documents Word (livrables 2 et 3)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ..models import ReportModel, Transcript
from . import common

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5A, 0x62, 0x6B)
WARNING = RGBColor(0x8A, 0x52, 0x00)

NOTHING_FOUND = "Aucun élément de ce type n'a été identifié dans cette réunion."


def _new_document() -> Document:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    return document


def _title_block(document: Document, title: str, subtitle: str) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT
    heading.paragraph_format.space_after = Pt(2)

    line = document.add_paragraph()
    run = line.add_run(subtitle)
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED
    line.paragraph_format.space_after = Pt(14)


def _section(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT


def _subsection(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def _label(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED


def _bullets(document: Document, lines: list[str]) -> None:
    for line in lines:
        if line and line.strip():
            document.add_paragraph(line.strip(), style="List Bullet")


def _note(document: Document, lines: list[str], *, color: RGBColor = MUTED) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = color


def _info_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    document.add_paragraph()


def _grid_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_paragraph()


# ---------------------------------------------------------------------------
# Livrable 2 — rapport complet
# ---------------------------------------------------------------------------


def render_report(model: ReportModel, transcript: Transcript | None, target: Path) -> Path:
    segments = common.segments_index(transcript)
    document = _new_document()

    _title_block(
        document,
        model.meta.title,
        f"Compte rendu de réunion — {model.meta.date} — durée {model.meta.duration_label}",
    )
    _note(document, [common.DISCLAIMER])

    # --- 1. Informations generales ---
    _section(document, "1. Informations générales")
    rows = [
        ("Titre", model.meta.title),
        ("Date", model.meta.date),
        ("Durée", model.meta.duration_label),
    ]
    if model.meta.subject:
        rows.append(("Objet", model.meta.subject))
    if model.meta.context:
        rows.append(("Contexte", model.meta.context))
    _info_table(document, rows)

    _label(document, "Participants identifiés")
    if model.meta.participants:
        _bullets(
            document,
            [
                " — ".join(
                    part
                    for part in (
                        p.name or p.label,
                        p.role,
                        f"({p.label})" if p.name else None,
                    )
                    if part
                )
                for p in model.meta.participants
            ],
        )
    else:
        document.add_paragraph(
            "Aucun participant n'a pu être identifié à partir de l'enregistrement."
        )

    if model.meta.missing_information:
        _label(document, "Informations manquantes ou non etablies")
        _bullets(document, model.meta.missing_information)

    quality = common.quality_lines(model)
    if quality:
        _label(document, "Limites de ce compte rendu")
        _note(document, quality, color=WARNING)

    # --- 2. Résumé exécutif ---
    _section(document, "2. Résumé exécutif")
    summary = model.executive_summary
    wrote_summary = False
    if summary.objectives:
        _label(document, "Objectifs")
        _bullets(document, summary.objectives)
        wrote_summary = True
    if summary.main_topics:
        _label(document, "Principaux sujets abordes")
        _bullets(document, summary.main_topics)
        wrote_summary = True
    if summary.key_outcomes:
        _label(document, "Résultats essentiels")
        _bullets(document, summary.key_outcomes)
        wrote_summary = True
    if not wrote_summary:
        document.add_paragraph(NOTHING_FOUND)

    # --- 3. Compte rendu détaillé ---
    _section(document, "3. Compte rendu détaillé")
    if not model.topics:
        document.add_paragraph(NOTHING_FOUND)
    for index, topic in enumerate(model.topics, start=1):
        _subsection(document, f"3.{index} {topic.title}")
        for paragraph in topic.narrative.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())

        if topic.key_points:
            _label(document, "Informations et données citées")
            _bullets(
                document,
                [kp.statement + common.reference(kp, segments) for kp in topic.key_points],
            )
        if topic.problems:
            _label(document, "Problèmes et causes évoquées")
            for problem in topic.problems:
                document.add_paragraph(
                    problem.statement + common.reference(problem, segments),
                    style="List Bullet",
                )
                for cause in problem.causes:
                    document.add_paragraph(f"Cause évoquée : {cause}", style="List Bullet 2")
        if topic.solutions:
            _label(document, "Solutions proposées et pistes étudiées")
            _bullets(
                document,
                [
                    f"{s.statement} ({common.SOLUTION_LABELS.get(s.status, s.status)})"
                    + common.reference(s, segments)
                    for s in topic.solutions
                ],
            )
        if topic.disagreements:
            _label(document, "Désaccords, réserves et questions soulevees")
            _bullets(
                document,
                [d.statement + common.reference(d, segments) for d in topic.disagreements],
            )
        if topic.related_topics:
            _note(document, ["En lien avec : " + ", ".join(topic.related_topics)])

    # --- 4. Décisions ---
    _section(document, "4. Décisions prises")
    if not model.decisions:
        document.add_paragraph(NOTHING_FOUND)
    for index, decision in enumerate(model.decisions, start=1):
        _subsection(
            document,
            f"4.{index} {decision.statement}{common.reference(decision, segments)}",
        )
        for name, value in common.decision_lines(decision, segments):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{name} : ")
            run.bold = True
            paragraph.add_run(value)

    # --- 5. Plan d'action ---
    _section(document, "5. Plan d'action")
    if model.actions:
        _grid_table(
            document,
            common.ACTION_HEADERS,
            [common.action_row(a, segments) for a in model.actions],
        )
        extras = [
            f"{a.task} — dépendances : {', '.join(a.dependencies + a.blockers)}"
            for a in model.actions
            if a.dependencies or a.blockers
        ]
        if extras:
            _label(document, "Dépendances et points bloquants")
            _bullets(document, extras)
    else:
        document.add_paragraph(NOTHING_FOUND)

    # --- 6. Questions ouvertes ---
    _section(document, "6. Questions ouvertes et points à clarifier")
    if model.open_questions:
        _bullets(
            document,
            [
                f"{q.question} ({common.QUESTION_LABELS.get(q.kind, q.kind)})" + common.reference(q, segments)
                for q in model.open_questions
            ],
        )
    else:
        document.add_paragraph(NOTHING_FOUND)

    if model.risks:
        _label(document, "Risques et points de vigilance")
        _bullets(
            document,
            [
                f"{r.statement} ({common.RISK_LABELS.get(r.kind, r.kind)})" + common.reference(r, segments)
                for r in model.risks
            ],
        )

    # --- 7. Conclusion ---
    _section(document, "7. Conclusion")
    if model.conclusion.assessment:
        document.add_paragraph(model.conclusion.assessment)
    if model.conclusion.next_steps:
        _label(document, "Prochaines étapes")
        _bullets(document, model.conclusion.next_steps)
    if model.conclusion.attention_points:
        _label(document, "Éléments nécessitant une attention particulière")
        _bullets(document, model.conclusion.attention_points)
    if not (
        model.conclusion.assessment
        or model.conclusion.next_steps
        or model.conclusion.attention_points
    ):
        document.add_paragraph(NOTHING_FOUND)

    _footer(document, model)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


# ---------------------------------------------------------------------------
# Livrable 3 — synthèse
# ---------------------------------------------------------------------------


def render_summary(model: ReportModel, transcript: Transcript | None, target: Path) -> Path:
    segments = common.segments_index(transcript)
    document = _new_document()

    _title_block(
        document,
        f"Synthèse — {model.meta.title}",
        f"{model.meta.date} — durée {model.meta.duration_label} — lecture en 2 minutes",
    )

    for title, lines in common.summary_sections(model, segments):
        _section(document, title)
        _bullets(document, lines)

    _section(document, "Actions à retenir")
    if model.actions:
        _grid_table(
            document,
            common.ACTION_HEADERS[:4],
            [common.action_row(a, segments)[:4] for a in model.actions],
        )
    else:
        document.add_paragraph("Aucune action n'a été identifiée.")

    quality = common.quality_lines(model)
    if quality:
        _section(document, "Fiabilité")
        _note(document, quality, color=WARNING)

    _footer(document, model)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _footer(document: Document, model: ReportModel) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(18)
    run = paragraph.add_run(
        f"Généré le {model.generated_at.strftime('%d/%m/%Y à %H:%M')} "
        f"— analyse : {model.llm_model or 'non précisée'}"
    )
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
